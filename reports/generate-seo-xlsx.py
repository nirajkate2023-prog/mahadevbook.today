from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# --- Checklist sheet ---
ws = wb.active
ws.title = "SEO Checklist"

headers = [
    "Priority",
    "Category",
    "Task",
    "Details",
    "Status",
    "Hours Spent",
    "Notes / Result",
    "Links",
]

rows = [
    ["HIGH", "Technical SEO", "Verify Google Search Console", "Make sure mahadevbooks.today is verified", "Done", 0.5, "Property verified in GSC (screenshots confirmed)", "https://search.google.com/search-console"],
    ["HIGH", "Technical SEO", "Create XML sitemap", "Create /sitemap.xml containing important public pages", "Done", 1, "17 URLs + lastmod dates added 8 Sep 2026", "https://mahadevbooks.today/sitemap.xml"],
    ["HIGH", "Technical SEO", "Submit sitemap", "Search Console → Sitemaps → submit sitemap.xml", "Done", 0.25, "Submitted Aug 24; Status Success; 17 pages discovered", ""],
    ["HIGH", "Technical SEO", "Create robots.txt", "Allow public pages and reference the sitemap; exclude private/admin areas", "Done", 0.5, "Allow / so Google can see noindex on login/register", "https://mahadevbooks.today/robots.txt"],
    ["HIGH", "Technical SEO", "Check indexing", "Search Console → Pages; investigate excluded/not-indexed important pages", "In Progress", 4, "4 indexed / 22 not indexed. Fixed 404/403/www/index.html/WP leftovers. Awaiting Google re-crawl.", ""],
    ["HIGH", "Technical SEO", "Check HTTPS", "Ensure all pages use HTTPS and HTTP redirects to HTTPS", "Done", 0.5, ".htaccess forces HTTPS + non-www", ""],
    ["HIGH", "Technical SEO", "Canonical URLs", "Add correct canonical URL to important indexable pages", "Done", 1.5, "Canonicals on all indexable pages; www→apex redirect added", ""],
    ["HIGH", "On-Page SEO", "Homepage title", "Use one clear natural title focused on the main topic/brand", "Done", 0.5, "Mahadev Book Official | Cricket ID & Sports Betting Platform", ""],
    ["HIGH", "On-Page SEO", "Homepage meta description", "Write a unique useful description matching the actual page", "Done", 0.5, "Unique description set in root layout", ""],
    ["HIGH", "On-Page SEO", "One H1 per important page", "Make the H1 clearly describe the page main topic", "Done", 1, "Single H1 pattern used across pages", ""],
    ["HIGH", "On-Page SEO", "Unique titles", "Every important page should have a unique title", "Done", 1.5, "Unique metadata.title per route (~19 pages)", ""],
    ["HIGH", "On-Page SEO", "Unique descriptions", "Write a unique meta description for each important page", "Done", 1.5, "Unique metadata.description per route", ""],
    ["MEDIUM", "On-Page SEO", "Heading structure", "Use H2/H3 headings logically under the H1", "Done", 1, "Guides/blogs use H1→H2 structure", ""],
    ["MEDIUM", "On-Page SEO", "Clean URLs", "Use short descriptive URLs", "Done", 0.5, "Kebab-case URLs with trailingSlash=true", ""],
    ["MEDIUM", "Images", "Image filenames", "Use descriptive filenames instead of image1.jpg", "Partial", 0.5, "Mix of descriptive and generic/WP filenames remain", ""],
    ["MEDIUM", "Images", "Alt text", "Add accurate alt text where images convey useful information", "Done", 1, "Descriptive alt on content images; avoid keyword stuffing", ""],
    ["HIGH", "Internal Linking", "Homepage links", "Link important sections from the homepage", "Partial", 1, "Footer quick links expanded (incl. mahadev-book-online). Homepage body still light.", ""],
    ["HIGH", "Internal Linking", "Topic clusters", "Link related guides and topic pages together", "Partial", 1.5, "Related links on several guides/blogs; needs stronger cluster linking", ""],
    ["MEDIUM", "Performance", "Page speed", "Test important pages with PageSpeed Insights", "Not Started", 0, "Static export helps; PSI report not run yet", "https://pagespeed.web.dev/"],
    ["MEDIUM", "Performance", "Core Web Vitals", "Check LCP INP and CLS in Search Console", "Not Started", 0, "Pending GSC Experience review", ""],
    ["HIGH", "Mobile SEO", "Mobile usability", "Check navigation buttons text images and scrolling on phones", "Partial", 1, "Responsive layout present; confirm in GSC Mobile Usability", ""],
    ["MEDIUM", "Structured Data", "Organization schema", "Add only if it accurately represents the organization", "Done", 1, "Organization + WebSite JSON-LD on homepage", ""],
    ["MEDIUM", "Structured Data", "Breadcrumb schema", "Add BreadcrumbList where breadcrumbs are present", "Done", 1, "BreadcrumbList via Breadcrumb component", ""],
    ["MEDIUM", "Structured Data", "Article schema", "Use on genuine blog/article pages", "Done", 1.5, "Article schema on guides and blog posts", ""],
    ["HIGH", "Content", "Create useful guides", "Publish genuinely helpful content answering real user questions", "Done", 12, "Cricket/IPL/UPI guides + 5 blog articles live (SEO content work)", ""],
    ["HIGH", "Content", "Avoid thin/duplicate content", "Do not create many pages with nearly identical text", "Partial", 1, "noindex login/register; WP leftovers redirected; topical overlap still a risk", ""],
    ["MEDIUM", "Authority", "Relevant mentions", "Pursue legitimate relevant websites that may reference useful content", "Not Started", 0, "No link-building campaign yet", ""],
    ["HIGH", "Monitoring", "Weekly Search Console review", "Track queries pages clicks impressions CTR and average position", "In Progress", 1, "Initial deep audit done 8 Sep 2026. Set weekly cadence going forward.", ""],
]

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(color="FFFFFF", bold=True)
thin = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
fills = {
    "Done": PatternFill("solid", fgColor="C6EFCE"),
    "In Progress": PatternFill("solid", fgColor="FFEB9C"),
    "Partial": PatternFill("solid", fgColor="FCE4D6"),
    "Not Started": PatternFill("solid", fgColor="FFC7CE"),
}

ws.append(headers)
for col in range(1, len(headers) + 1):
    cell = ws.cell(1, col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center")

for row in rows:
    ws.append(row)
    r = ws.max_row
    status = row[4]
    for col in range(1, len(headers) + 1):
        cell = ws.cell(r, col)
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if col == 5:
            cell.fill = fills.get(status, PatternFill())
            cell.font = Font(bold=True)

widths = [10, 16, 28, 55, 14, 12, 55, 40]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 24
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:H{ws.max_row}"

# --- Summary sheet ---
ws2 = wb.create_sheet("Time Summary")
total = sum(r[5] for r in rows)
counts = {k: sum(1 for r in rows if r[4] == k) for k in ["Done", "In Progress", "Partial", "Not Started"]}
summary = [
    ["Metric", "Value"],
    ["Report date", "8 Sep 2026"],
    ["Domain", "mahadevbooks.today"],
    ["SEO status", "Technical setup mostly Done; indexing recovery In Progress; traffic still very low"],
    ["Indexed pages (GSC)", 4],
    ["Not indexed pages (GSC)", 22],
    ["Clicks (last ~3 months)", 1],
    ["Impressions (last ~3 months)", 15],
    ["Sitemap", "Success - 17 pages discovered"],
    ["Tasks Done", counts["Done"]],
    ["Tasks In Progress", counts["In Progress"]],
    ["Tasks Partial", counts["Partial"]],
    ["Tasks Not Started", counts["Not Started"]],
    ["Total SEO hours spent", total],
    ["Hours for client report", f"{total} hrs (range 35-40 hrs)"],
    ["Next focus", "Request indexing for main URLs; weekly GSC check; CWV/speed; stronger internal links"],
]
for row in summary:
    ws2.append(row)
ws2["A1"].fill = header_fill
ws2["B1"].fill = header_fill
ws2["A1"].font = header_font
ws2["B1"].font = header_font
ws2.column_dimensions["A"].width = 28
ws2.column_dimensions["B"].width = 90

# --- Hours by category ---
ws3 = wb.create_sheet("Hours By Category")
ws3.append(["Category", "Hours", "Tasks", "Done"])
for col in range(1, 5):
    ws3.cell(1, col).fill = header_fill
    ws3.cell(1, col).font = header_font

by_cat = {}
for r in rows:
    cat = r[1]
    cur = by_cat.setdefault(cat, {"hours": 0, "tasks": 0, "done": 0})
    cur["hours"] += r[5]
    cur["tasks"] += 1
    if r[4] == "Done":
        cur["done"] += 1

for cat, v in sorted(by_cat.items(), key=lambda x: x[1]["hours"], reverse=True):
    ws3.append([cat, v["hours"], v["tasks"], v["done"]])

for col, w in enumerate([18, 10, 10, 10], 1):
    ws3.column_dimensions[get_column_letter(col)].width = w

out = r"d:\Downloads\mahadevbook.today\reports\Mahadev_Books_SEO_Report_Updated.xlsx"
wb.save(out)
print(out)
print(f"Total hours: {total}")
